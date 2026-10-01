"""Read weekly sales files, isolate exceptions, and write an auditable workbook."""

from __future__ import annotations

import argparse
import csv
import os
import re
import tempfile
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile
from zlib import error as ZlibError

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils.exceptions import InvalidFileException

from .presentation import build_workbook


REQUIRED = (
    "sale_date", "branch_code", "order_id", "line_no", "product_code",
    "category", "quantity", "unit_price", "status",
)
METRICS = (
    "discovered_input_file_count", "processed_file_count", "rejected_file_count",
    "total_source_row_count", "clean_row_count", "row_exception_count",
    "rejected_file_source_row_count", "file_exception_count", "duplicate_row_count",
    "total_sales_amount",
)


class SchemaError(ValueError):
    def __init__(self, message: str, rows: int | None = None):
        super().__init__(message)
        self.rows = rows


def _headers(raw_headers: tuple) -> tuple[str, ...]:
    names = [str(value).strip().lower() if value is not None else "" for value in raw_headers]
    while names and not names[-1]:
        names.pop()
    nonempty = [name for name in names if name]
    if len(nonempty) != len(set(nonempty)):
        raise SchemaError("Duplicate column names after trim/lowercase normalization")
    if "" in names:
        raise SchemaError("Unnamed column within the header; name it or remove its data")
    if any(ILLEGAL_CHARACTERS_RE.search(name) for name in names):
        raise SchemaError("Header contains a character Excel cannot write")
    missing = [name for name in REQUIRED if name not in names]
    if missing:
        raise SchemaError("Missing required columns: " + ", ".join(missing))
    return tuple(names)


def _blank(values) -> bool:
    return all(value is None or isinstance(value, str) and not value.strip() for value in values)


def _csv_records(path: Path) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            with path.open("r", encoding=encoding, newline="") as stream:
                return list(csv.reader(stream, strict=True))
        except UnicodeDecodeError:
            continue
    raise UnicodeError("CSV encoding is unsupported; save as UTF-8 (with or without BOM) or CP949")


def _read_file(path: Path) -> tuple[list[tuple[int, dict, tuple[str, ...]]], int]:
    if path.suffix.lower() == ".csv":
        records = _csv_records(path)
        if not records:
            raise SchemaError("Empty file", 0)
        source_rows = sum(not _blank(record) for record in records[1:])
        try:
            names = _headers(tuple(records[0]))
        except SchemaError as exc:
            raise SchemaError(str(exc), source_rows) from exc
        rows = []
        for row_number, values in enumerate(records[1:], start=2):
            if _blank(values):
                continue
            if len(values) < len(names) or not _blank(values[len(names):]):
                raise SchemaError(f"Row {row_number} has {len(values)} cells; expected {len(names)}", source_rows)
            rows.append((row_number, dict(zip(names, values[:len(names)])), ()))
        return rows, source_rows

    # Inspect formulas themselves: cached results can be missing or stale.
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        # Read-only iteration otherwise trusts potentially stale <dimension> metadata.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        header = next(iterator, None)
        if header is None:
            raise SchemaError("Empty file", 0)
        values = list(iterator)
        source_rows = sum(not _blank(cell.value for cell in cells) for cells in values)
        try:
            names = _headers(tuple(cell.value for cell in header))
        except SchemaError as exc:
            raise SchemaError(str(exc), source_rows) from exc
        rows = []
        for row_number, cells in enumerate(values, start=2):
            cell_values = tuple(cell.value for cell in cells)
            if _blank(cell_values):
                continue
            if len(cells) > len(names) and any(value is not None for value in cell_values[len(names):]):
                raise SchemaError(f"Row {row_number} has data beyond the header columns", source_rows)
            # Sparse XLSX rows may end before the header width. Keep missing
            # values explicit so the normal row validator can report them.
            raw = {name: cell_values[index] if index < len(cell_values) else None
                   for index, name in enumerate(names)}
            formula_columns = tuple(name for name, cell in zip(names, cells) if cell.data_type == "f")
            for name in formula_columns:
                value = raw[name]
                if not isinstance(value, str):
                    # Array/data-table formulas have structured values in openpyxl.
                    raw[name] = getattr(value, "text", None) or f"{type(value).__name__}: {vars(value)}"
            rows.append((row_number, raw, formula_columns))
        return rows, source_rows
    finally:
        workbook.close()


def _string(value: object) -> str:
    return "" if value is None else str(value).strip()


def _identifier(value: object) -> str | None:
    if isinstance(value, str):
        return value.strip() or None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        return None
    if (not number.is_finite() or number < 0 or number != number.to_integral_value()
            or len(str(int(number))) > 15):
        return None
    return str(int(number))


def _excel_safe(value: object) -> object:
    if not isinstance(value, str):
        return value
    return ILLEGAL_CHARACTERS_RE.sub(lambda match: f"\\x{ord(match.group()):02X}", value)


def _positive_integer(value: object) -> int | None:
    try:
        number = Decimal(_string(value))
    except InvalidOperation:
        return None
    if not number.is_finite() or number <= 0 or number != number.to_integral_value():
        return None
    return int(number)


def _price(value: object) -> Decimal | None:
    try:
        number = Decimal(_string(value))
    except InvalidOperation:
        return None
    return number if number.is_finite() and number >= 0 else None


def _date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _string(value)
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _validate(raw: dict) -> tuple[dict, list[str]]:
    errors = []
    sale_date = _date(raw["sale_date"])
    if sale_date is None:
        errors.append("INVALID_DATE")
    strings = {}
    for field in ("branch_code", "order_id", "product_code", "category", "status"):
        original = raw[field]
        value = _identifier(original) if field in ("branch_code", "order_id", "product_code") else _string(original)
        if original is None or isinstance(original, str) and not original.strip():
            errors.append(f"MISSING_{field.upper()}")
        elif value is None or field in ("category", "status") and not isinstance(original, str):
            errors.append(f"INVALID_{field.upper()}")
        value = value or ""
        strings[field] = " ".join(value.upper().split()) if field == "status" else value
    line_no = _positive_integer(raw["line_no"])
    if line_no is None:
        errors.append("INVALID_LINE_NO")
    quantity = _positive_integer(raw["quantity"])
    if quantity is None:
        errors.append("INVALID_QUANTITY")
    unit_price = _price(raw["unit_price"])
    if unit_price is None:
        errors.append("INVALID_UNIT_PRICE")
    if strings["status"] and strings["status"] != "COMPLETED":
        errors.append("UNSUPPORTED_STATUS")
    if any(isinstance(value, str) and ILLEGAL_CHARACTERS_RE.search(value) for value in raw.values()):
        errors.append("EXCEL_ILLEGAL_CHARACTER")
    normalized = {
        "sale_date": sale_date, "branch_code": strings["branch_code"],
        "order_id": strings["order_id"], "line_no": line_no,
        "product_code": strings["product_code"], "category": strings["category"],
        "quantity": quantity, "unit_price": unit_price, "status": strings["status"],
    }
    return normalized, errors


def _write_report(path: Path, clean: list[dict], exceptions: list[dict], runs: list[dict],
                  original_columns: list[str]) -> dict:
    branch = defaultdict(Decimal)
    category = defaultdict(Decimal)
    for row in clean:
        branch[row["branch_code"]] += row["sale_amount"]
        category[row["category"]] += row["sale_amount"]
    metrics = {
        "discovered_input_file_count": len(runs),
        "processed_file_count": sum(row["status"] == "PROCESSED" for row in runs),
        "rejected_file_count": sum(row["status"] == "REJECTED" for row in runs),
        "total_source_row_count": sum(row["source_row_count"] or 0 for row in runs),
        "clean_row_count": len(clean),
        "row_exception_count": sum(row["source_row"] is not None for row in exceptions),
        "rejected_file_source_row_count": sum(
            row["source_row_count"] or 0 for row in runs if row["status"] == "REJECTED"
        ),
        "file_exception_count": sum(row["source_row"] is None for row in exceptions),
        "duplicate_row_count": sum("DUPLICATE_TRANSACTION" in row["error_codes"].split("|") for row in exceptions),
        "total_sales_amount": sum((row["sale_amount"] for row in clean), Decimal(0)),
    }
    workbook = build_workbook(clean, exceptions, runs, original_columns, metrics, branch, category, REQUIRED)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".xlsx", prefix=".weekly_sales_", dir=path.parent, delete=False) as stream:
            temporary = Path(stream.name)
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return metrics


def build_report(input_dir: Path, output_path: Path) -> dict:
    input_dir = input_dir.resolve()
    output_path = output_path.resolve()
    if not input_dir.is_dir():
        raise ValueError(f"Input directory does not exist: {input_dir}")
    if output_path.parent == input_dir:
        raise ValueError("Output report must be outside the input directory")
    files = sorted(
        (path for path in input_dir.iterdir()
         if path.is_file() and path.suffix.lower() in (".csv", ".xlsx", ".xls", ".xlsm", ".xlsb")
         and not path.name.startswith("~$") and path.resolve() != output_path),
        key=lambda path: path.name,
    )
    if not files:
        raise ValueError(f"No CSV or XLSX input files in: {input_dir}")
    clean, exceptions, runs = [], [], []
    original_columns = list(REQUIRED)
    seen = {}
    for path in files:
        run = {"filename": path.name, "format": path.suffix.lower(), "status": "PROCESSED",
               "source_row_count": None, "clean_contribution": 0,
               "exception_contribution": 0, "error_message": ""}
        if path.suffix.lower() not in (".csv", ".xlsx"):
            message = "Unsupported spreadsheet type; save as CSV or XLSX"
            run.update(status="REJECTED", exception_contribution=1, error_message=message)
            exceptions.append({"source_file": path.name, "source_row": None,
                               "error_codes": "UNSUPPORTED_FILE_FORMAT", "error_detail": message})
            runs.append(run)
            continue
        try:
            rows, count = _read_file(path)
            run["source_row_count"] = count
        except SchemaError as exc:
            run.update(status="REJECTED", source_row_count=exc.rows,
                       exception_contribution=1, error_message=str(exc))
            code = "MISSING_REQUIRED_COLUMNS" if str(exc).startswith("Missing required columns") else "FILE_SCHEMA_ERROR"
            exceptions.append({"source_file": path.name, "source_row": None,
                               "error_codes": code, "error_detail": str(exc)})
            runs.append(run)
            continue
        except (OSError, UnicodeError, ValueError, KeyError, IndexError, BadZipFile,
                ParseError, ZlibError, EOFError, InvalidFileException, csv.Error) as exc:
            if isinstance(exc, UnicodeError) and "CSV encoding is unsupported" in str(exc):
                message = str(exc)
            else:
                message = f"Unreadable input: {type(exc).__name__}: {exc}"
            run.update(status="REJECTED", exception_contribution=1, error_message=message)
            exceptions.append({"source_file": path.name, "source_row": None,
                               "error_codes": "FILE_READ_ERROR", "error_detail": message})
            runs.append(run)
            continue
        for source_row, raw, formula_columns in rows:
            for name in raw:
                if name not in original_columns:
                    original_columns.append(name)
            normalized, errors = ({}, ["UNSUPPORTED_FORMULA"]) if formula_columns else _validate(raw)
            if not errors:
                key = (normalized["order_id"], normalized["line_no"])
                if key in seen:
                    errors.append("DUPLICATE_TRANSACTION")
                else:
                    seen[key] = normalized.copy(), path.name, source_row
            if errors:
                detail = "Row failed validation"
                if formula_columns:
                    detail = ("XLSX formulas are not supported; export values before processing. Formula columns: "
                              + ", ".join(formula_columns))
                elif errors == ["DUPLICATE_TRANSACTION"]:
                    first, first_file, first_row = seen[key]
                    different = [name for name in REQUIRED if normalized[name] != first[name]]
                    detail = (f"Conflicts with first accepted row {first_file}:{first_row}; different fields: "
                              + ", ".join(different)) if different else "Duplicate accepted transaction"
                elif "EXCEL_ILLEGAL_CHARACTER" in errors:
                    detail = "Customer text contains Excel-illegal control characters; affected cells show escaped codes"
                exception = {"source_file": path.name, "source_row": source_row,
                             "error_codes": "|".join(errors),
                             "error_detail": detail}
                exception.update({"original_" + name: _excel_safe(value) for name, value in raw.items()})
                exceptions.append(exception)
                run["exception_contribution"] += 1
            else:
                normalized.update(source_file=path.name, source_row=source_row,
                                  sale_amount=normalized["quantity"] * normalized["unit_price"])
                clean.append(normalized)
                run["clean_contribution"] += 1
        runs.append(run)
    return _write_report(output_path, clean, exceptions, runs, original_columns)


def main() -> int:
    parser = argparse.ArgumentParser(description="Consolidate weekly CSV/XLSX sales into an Excel report")
    parser.add_argument("--input", required=True, type=Path, help="Directory of CSV/XLSX sales files")
    parser.add_argument("--output", required=True, type=Path, help="Output .xlsx report path")
    args = parser.parse_args()
    if args.output.suffix.lower() != ".xlsx":
        parser.error("--output must end in .xlsx")
    try:
        metrics = build_report(args.input, args.output)
    except (OSError, ValueError) as exc:
        parser.exit(1, f"Error: {exc}\n")
    print(f"Report: {args.output.resolve()}")
    for name in METRICS:
        print(f"{name}: {metrics[name]}")
    return 0
