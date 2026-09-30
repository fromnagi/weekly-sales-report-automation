"""Workbook presentation for the already-validated sales results."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
import math

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "20364B"
TEAL = "0C7C86"
TEXT = "243746"
MUTED = "607181"
PALE = "F3F6F8"
PALE_BLUE = "EAF2F5"
PALE_AMBER = "FFF3DC"
PALE_RED = "FBEAEA"
RED = "9A3438"
GREEN = "32685A"
WHITE = "FFFFFF"
AMOUNT_FORMAT = "#,##0.00"
COUNT_FORMAT = "#,##0"

RUN_COLUMNS = (
    "filename", "format", "status", "source_row_count", "clean_contribution",
    "exception_contribution", "error_message",
)
EXCEPTION_COLUMNS = (
    "error_codes", "error_detail", "failure_scope", "source_file", "source_row",
)
AUDIT_LABELS = {
    "discovered_input_file_count": "Files received",
    "processed_file_count": "Files processed",
    "rejected_file_count": "Files rejected",
    "total_source_row_count": "Readable source rows",
    "clean_row_count": "Accepted rows",
    "row_exception_count": "Rows to review",
    "rejected_file_source_row_count": "Rows in rejected files",
    "file_exception_count": "File exception records",
    "duplicate_row_count": "Duplicate rows",
    "total_sales_amount": "Total sales amount",
}


def _value(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def _font(*, size=10, color=TEXT, bold=False):
    return Font(name="Aptos", size=size, color=color, bold=bold)


def _band(sheet, cell_range: str, text: str) -> None:
    sheet.merge_cells(cell_range)
    cell = sheet[cell_range.split(":")[0]]
    cell.value = text
    cell.font = _font(size=10, color=WHITE, bold=True)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[cell.row].height = 25


def _card(sheet, start: str, end: str, label: str, value, *, amount=False, warning=False) -> None:
    first_col = sheet[start].column_letter
    last_col = sheet[end].column_letter
    top = sheet[start].row
    bottom = sheet[end].row
    for row in sheet.iter_rows(min_row=top, max_row=bottom,
                               min_col=sheet[start].column, max_col=sheet[end].column):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=PALE_AMBER if warning else PALE_BLUE)
    sheet.merge_cells(f"{first_col}{top}:{last_col}{top}")
    sheet.merge_cells(f"{first_col}{top + 1}:{last_col}{bottom}")
    label_cell = sheet[start]
    label_cell.value = label
    label_cell.font = _font(size=10, color=MUTED, bold=True)
    label_cell.alignment = Alignment(horizontal="center", vertical="center")
    number = sheet[f"{first_col}{top + 1}"]
    number.value = _value(value)
    number.font = _font(size=19, color=RED if warning else NAVY, bold=True)
    number.number_format = AMOUNT_FORMAT if amount else COUNT_FORMAT
    number.alignment = Alignment(horizontal="center", vertical="center")


def _chart(sheet, *, first_col: int, last_col: int, count: int, title: str,
           anchor: str, color: str) -> None:
    if count == 0:
        start = anchor[0]
        end = "H" if start == "B" else "P"
        sheet.merge_cells(f"{start}17:{end}19")
        cell = sheet[f"{start}17"]
        cell.value = "No accepted sales to chart"
        cell.font = _font(size=11, color=MUTED)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        return
    chart = BarChart()
    chart.type = "bar"
    chart.title = title
    chart.legend = None
    chart.layout = Layout(manualLayout=ManualLayout(x=0.26, y=0.14, w=0.70, h=0.70))
    axis_text = RichText(p=[Paragraph(
        pPr=ParagraphProperties(defRPr=CharacterProperties(solidFill=TEXT, sz=900))
    )])
    chart.x_axis.axPos = "l"
    chart.x_axis.tickLblPos = "nextTo"
    chart.x_axis.txPr = axis_text
    chart.y_axis.axPos = "b"
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.txPr = axis_text
    chart.y_axis.numFmt = AMOUNT_FORMAT
    chart.width = 13.1
    chart.height = 7.2
    chart.add_data(Reference(sheet, min_col=last_col, min_row=29, max_row=29 + count),
                   titles_from_data=True)
    category_col = get_column_letter(first_col)
    chart.series[0].cat = AxDataSource(
        strRef=StrRef(f=f"'{sheet.title}'!${category_col}$30:${category_col}${29 + count}")
    )
    chart.dLbls = DataLabelList()
    chart.dLbls.showCatName = True
    chart.dLbls.showSerName = False
    chart.dLbls.showVal = False
    chart.dLbls.dLblPos = "inBase"
    chart.dLbls.txPr = RichText(p=[Paragraph(
        pPr=ParagraphProperties(defRPr=CharacterProperties(solidFill=WHITE, sz=900))
    )])
    chart.series[0].graphicalProperties.solidFill = color
    chart.series[0].graphicalProperties.line.noFill = True
    sheet.add_chart(chart, anchor)


def _aggregate_table(sheet, col: int, label: str, values: dict, color: str) -> None:
    left = sheet.cell(29, col, label)
    right = sheet.cell(29, col + 1, "Sales amount")
    for cell in (left, right):
        cell.font = _font(color=WHITE, bold=True)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[29].height = 24
    for offset, (name, amount) in enumerate(sorted(values.items()), start=30):
        name_cell = sheet.cell(offset, col, name)
        name_cell.data_type = "s"
        amount_cell = sheet.cell(offset, col + 1, float(amount))
        for cell in (name_cell, amount_cell):
            cell.font = _font()
            cell.fill = PatternFill("solid", fgColor=WHITE if offset % 2 else PALE)
            cell.alignment = Alignment(vertical="center", indent=1)
        amount_cell.number_format = AMOUNT_FORMAT
        sheet.row_dimensions[offset].height = 21


def _health_item(sheet, row: int, label_col: str, value_col: str, label: str, value: int) -> None:
    label_start = sheet[f"{label_col}{row}"]
    sheet.merge_cells(start_row=row, start_column=label_start.column,
                      end_row=row, end_column=sheet[f"{value_col}{row}"].column - 1)
    label_start.value = label
    label_start.font = _font(color=MUTED)
    label_start.alignment = Alignment(vertical="center", indent=1)
    value_cell = sheet[f"{value_col}{row}"]
    value_cell.value = value
    value_cell.font = _font(color=NAVY, bold=True)
    value_cell.number_format = COUNT_FORMAT
    value_cell.alignment = Alignment(horizontal="right", vertical="center")


def _summary(sheet, clean: list[dict], metrics: dict, branch: dict, category: dict) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.sheet_properties.tabColor = NAVY
    sheet.column_dimensions["A"].width = 3
    sheet.column_dimensions["Q"].width = 3
    for col in range(2, 17):
        sheet.column_dimensions[get_column_letter(col)].width = 10.5
    for col in ("B", "J"):
        sheet.column_dimensions[col].width = 12.5
    for col in ("C", "K"):
        sheet.column_dimensions[col].width = 13.5

    sheet.merge_cells("B2:P2")
    sheet["B2"] = "WEEKLY SALES REPORT"
    sheet["B2"].font = _font(size=20, color=NAVY, bold=True)
    sheet.row_dimensions[2].height = 34
    sheet.merge_cells("B3:P3")
    sheet["B3"] = "Automated multi-branch sales consolidation"
    sheet["B3"].font = _font(size=11, color=MUTED)
    sheet.row_dimensions[3].height = 22
    sheet.merge_cells("B4:P4")
    if clean:
        dates = [row["sale_date"] for row in clean]
        period = f"{min(dates):%d %b %Y}  –  {max(dates):%d %b %Y}"
    else:
        period = "No accepted sales dates"
    sheet["B4"] = f"Reporting period: {period}"
    sheet["B4"].font = _font(size=10, color=MUTED)
    sheet.row_dimensions[4].height = 21
    for col in range(2, 17):
        sheet.cell(5, col).border = Border(bottom=Side(style="medium", color=TEAL))

    _card(sheet, "B7", "D10", "TOTAL SALES", metrics["total_sales_amount"], amount=True)
    _card(sheet, "F7", "H10", "ACCEPTED ROWS", metrics["clean_row_count"])
    _card(sheet, "J7", "L10", "ROWS TO REVIEW", metrics["row_exception_count"], warning=True)
    _card(sheet, "N7", "P10", "REJECTED FILES", metrics["rejected_file_count"], warning=True)
    sheet.row_dimensions[7].height = 25
    for row in (8, 9, 10):
        sheet.row_dimensions[row].height = 21

    _aggregate_table(sheet, 2, "Branch", branch, NAVY)
    _aggregate_table(sheet, 10, "Category", category, TEAL)
    _chart(sheet, first_col=2, last_col=3, count=len(branch),
           title="Sales by Branch", anchor="B13", color=TEAL)
    _chart(sheet, first_col=10, last_col=11, count=len(category),
           title="Sales by Category", anchor="J13", color=NAVY)

    health = max(35, 31 + max(len(branch), len(category)))
    _band(sheet, f"B{health}:P{health}", "PROCESSING HEALTH")
    health_rows = (
        (("Files received", "discovered_input_file_count"),
         ("Processed", "processed_file_count"), ("Rejected", "rejected_file_count")),
        (("Readable source rows", "total_source_row_count"),
         ("Accepted", "clean_row_count"), ("Review rows", "row_exception_count")),
        (("Rejected-file rows", "rejected_file_source_row_count"),
         ("Duplicates", "duplicate_row_count"), ("File exception records", "file_exception_count")),
    )
    for offset, items in enumerate(health_rows, start=2):
        row = health + offset
        sheet.row_dimensions[row].height = 24
        for (label, key), (label_col, value_col) in zip(items, (("B", "E"), ("G", "J"), ("L", "O"))):
            _health_item(sheet, row, label_col, value_col, label, metrics[key])
    reconcile_row = health + 5
    sheet.merge_cells(start_row=reconcile_row, start_column=2, end_row=reconcile_row, end_column=16)
    sheet.cell(reconcile_row, 2, (
        f"Source rows reconcile: {metrics['total_source_row_count']:,} = "
        f"{metrics['clean_row_count']:,} accepted + {metrics['row_exception_count']:,} review + "
        f"{metrics['rejected_file_source_row_count']:,} in rejected files"
    ))
    sheet.cell(reconcile_row, 2).font = _font(color=TEAL, bold=True)
    sheet.row_dimensions[reconcile_row].height = 27

    audit = health + 7
    _band(sheet, f"B{audit}:P{audit}", "AUDIT METRICS")
    sheet.merge_cells(start_row=audit + 1, start_column=2, end_row=audit + 1, end_column=8)
    sheet.merge_cells(start_row=audit + 1, start_column=10, end_row=audit + 1, end_column=16)
    for col, title in ((2, "Measure"), (9, "Value"), (10, "Metric key")):
        cell = sheet.cell(audit + 1, col, title)
        cell.font = _font(color=MUTED, bold=True)
    sheet.row_dimensions[audit + 1].height = 23
    for offset, (key, value) in enumerate(metrics.items(), start=2):
        row = audit + offset
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        sheet.merge_cells(start_row=row, start_column=10, end_row=row, end_column=16)
        for col in range(2, 17):
            sheet.cell(row, col).fill = PatternFill("solid", fgColor=PALE if offset % 2 == 0 else WHITE)
        label = sheet.cell(row, 2, AUDIT_LABELS[key])
        label.font = _font()
        label.alignment = Alignment(vertical="center", indent=1)
        number = sheet.cell(row, 9, _value(value))
        number.font = _font(color=NAVY, bold=True)
        number.number_format = AMOUNT_FORMAT if key == "total_sales_amount" else COUNT_FORMAT
        number.alignment = Alignment(horizontal="right", vertical="center")
        code = sheet.cell(row, 10, key)
        code.font = _font(size=9, color=MUTED)
        code.alignment = Alignment(vertical="center", indent=1)
        sheet.row_dimensions[row].height = 21

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4,
                                     header=0.15, footer=0.15)
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f"B2:P{audit + len(metrics) + 3}"


def _detail(sheet, columns: tuple[str, ...], records: list[dict], widths: dict[str, int],
            *, tab_color: str, freeze: str, table_name: str | None = None) -> None:
    sheet.append(columns)
    for record in records:
        sheet.append([_value(record.get(column)) for column in columns])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_properties.tabColor = tab_color
    sheet.freeze_panes = freeze
    if not (table_name and records):
        sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.font = _font(color=WHITE, bold=True)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 23
        for cell in row:
            cell.font = _font()
            cell.fill = PatternFill("solid", fgColor=WHITE if cell.row % 2 else PALE)
            cell.alignment = Alignment(vertical="center", indent=1,
                                       horizontal="right" if isinstance(cell.value, (int, float, date)) else "left")
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(column, 19)
    if table_name and records:
        table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(columns))}{len(records) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                             showLastColumn=False, showRowStripes=True,
                                             showColumnStripes=False)
        sheet.add_table(table)


def build_workbook(clean: list[dict], exceptions: list[dict], runs: list[dict],
                   original_columns: list[str], metrics: dict, branch: dict, category: dict,
                   required: tuple[str, ...]) -> Workbook:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _summary(summary, clean, metrics, branch, category)

    clean_columns = (*required, "source_file", "source_row", "sale_amount")
    clean_sheet = workbook.create_sheet("Clean_Data")
    _detail(clean_sheet, clean_columns, clean, {
        "sale_date": 16, "branch_code": 17, "order_id": 18, "line_no": 12,
        "product_code": 20, "category": 19, "quantity": 13, "unit_price": 16,
        "status": 17, "source_file": 24, "source_row": 14, "sale_amount": 18,
    }, tab_color=TEAL, freeze="D2", table_name="CleanSales")
    for row in clean_sheet.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        for index in (3, 6, 10):
            row[index].number_format = COUNT_FORMAT
        for index in (7, 11):
            row[index].number_format = AMOUNT_FORMAT

    exception_columns = (*EXCEPTION_COLUMNS, *("original_" + name for name in original_columns))
    display_exceptions = [dict(record, failure_scope="FILE" if record["source_row"] is None else "ROW")
                          for record in exceptions]
    exception_sheet = workbook.create_sheet("Exceptions")
    _detail(exception_sheet, exception_columns, display_exceptions, {
        "error_codes": 38, "error_detail": 48, "failure_scope": 16,
        "source_file": 25, "source_row": 14,
    }, tab_color="B67A24", freeze="D2")
    for row in exception_sheet.iter_rows(min_row=2):
        is_file = row[2].value == "FILE"
        row[0].font = _font(color=RED, bold=True)
        row[0].fill = PatternFill("solid", fgColor=PALE_RED if is_file else PALE_AMBER)
        row[0].alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        row[1].alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        row[2].font = _font(color=RED if is_file else TEXT, bold=True)
        row[2].fill = PatternFill("solid", fgColor=PALE_RED if is_file else PALE_AMBER)
        exception_sheet.row_dimensions[row[0].row].height = 39 if is_file else 34

    run_sheet = workbook.create_sheet("Run_Info")
    _detail(run_sheet, RUN_COLUMNS, runs, {
        "filename": 27, "format": 12, "status": 17, "source_row_count": 19,
        "clean_contribution": 21, "exception_contribution": 23, "error_message": 52,
    }, tab_color=MUTED, freeze="D2")
    for row in run_sheet.iter_rows(min_row=2):
        rejected = row[2].value == "REJECTED"
        row[2].font = _font(color=RED if rejected else GREEN, bold=True)
        row[2].fill = PatternFill("solid", fgColor=PALE_RED if rejected else "E8F3EF")
        row[6].alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        run_sheet.row_dimensions[row[0].row].height = 37 if row[6].value else 24
        for index in (3, 4, 5):
            row[index].number_format = COUNT_FORMAT
    return workbook
