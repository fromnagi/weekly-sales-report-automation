import csv
import hashlib
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from weekly_sales_report.report import _identifier, build_report


HEADERS = [
    "sale_date", "branch_code", "order_id", "line_no", "product_code",
    "category", "quantity", "unit_price", "status",
]


def write_csv(path: Path, headers: list, rows: list[list]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)


def write_xlsx(path: Path, headers: list, rows: list[list]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def rows(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def hash_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def snapshot(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {sheet.title: list(sheet.values) for sheet in workbook}
    finally:
        workbook.close()


def summary_metrics(sheet) -> dict:
    return {sheet.cell(row, 10).value: sheet.cell(row, 9).value
            for row in range(1, sheet.max_row + 1)
            if sheet.cell(row, 10).value in {
                "discovered_input_file_count", "processed_file_count", "rejected_file_count",
                "total_source_row_count", "clean_row_count", "row_exception_count",
                "rejected_file_source_row_count", "file_exception_count",
                "duplicate_row_count", "total_sales_amount",
            }}


def test_mixed_files_reconcile_preserve_bytes_and_rerun(tmp_path):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    output = tmp_path / "weekly_sales_report.xlsx"
    write_csv(input_dir / "01_sales.csv", [f" {name.upper()} " for name in HEADERS], [
        ["2026-09-21", "B1", "O1", "1", "P1", "Food", "2", "10.50", " completed "],
        ["2026-09-21", "B1", "O2", "1", "P2", "Home", "1", "5", "COMPLETED"],
        ["bad", "B1", "O3", "1", "", "Food", "x", "2", "COMPLETED"],
    ])
    write_xlsx(input_dir / "02_sales.xlsx", HEADERS, [
        ["2026-09-21", "B2", "O1", 1, "P9", "Food", 7, 100, "COMPLETED"],
        ["2026-09-22", "B2", "O4", 2, "P3", "Food", 3, 4, "COMPLETED"],
    ])
    write_csv(input_dir / "03_missing.csv", [name for name in HEADERS if name != "quantity"], [
        ["2026-09-21", "B3", "O8", "1", "P4", "Home", "900", "COMPLETED"],
        ["2026-09-21", "B3", "O9", "1", "P5", "Home", "900", "COMPLETED"],
    ])
    before = {path.name: hash_file(path) for path in input_dir.iterdir()}
    expected = {
        "discovered_input_file_count": 3, "processed_file_count": 2,
        "rejected_file_count": 1, "total_source_row_count": 7,
        "clean_row_count": 3, "row_exception_count": 2,
        "rejected_file_source_row_count": 2, "file_exception_count": 1,
        "duplicate_row_count": 1, "total_sales_amount": 38,
    }
    assert expected["total_source_row_count"] == (
        expected["clean_row_count"] + expected["row_exception_count"]
        + expected["rejected_file_source_row_count"]
    )
    assert build_report(input_dir, output) == expected
    assert {path.name: hash_file(path) for path in input_dir.iterdir()} == before
    first = snapshot(output)
    assert list(first) == ["Summary", "Clean_Data", "Exceptions", "Run_Info"]
    assert build_report(input_dir, output) == expected
    assert snapshot(output) == first
    assert {path.name: hash_file(path) for path in input_dir.iterdir()} == before

    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        clean = rows(workbook["Clean_Data"])
        exceptions = rows(workbook["Exceptions"])
        runs = rows(workbook["Run_Info"])
        summary = workbook["Summary"]
        assert [(row["order_id"], row["source_file"], row["source_row"]) for row in clean] == [
            ("O1", "01_sales.csv", 2), ("O2", "01_sales.csv", 3),
            ("O4", "02_sales.xlsx", 3),
        ]
        assert [row["sale_amount"] for row in clean] == [21, 5, 12]
        assert exceptions[0]["error_codes"] == "INVALID_DATE|MISSING_PRODUCT_CODE|INVALID_QUANTITY"
        assert exceptions[1]["error_codes"] == "DUPLICATE_TRANSACTION"
        assert exceptions[2]["error_codes"] == "MISSING_REQUIRED_COLUMNS"
        assert exceptions[2]["source_row"] is None
        assert [row["status"] for row in runs] == ["PROCESSED", "PROCESSED", "REJECTED"]
        assert [row["clean_contribution"] for row in runs] == [2, 1, 0]
        assert [row["exception_contribution"] for row in runs] == [1, 1, 1]
        assert summary_metrics(summary) == expected
        assert summary["B2"].value == "WEEKLY SALES REPORT"
        assert "21 Sep 2026" in summary["B4"].value
        assert "22 Sep 2026" in summary["B4"].value
        assert [summary[cell].value for cell in ("B8", "F8", "J8", "N8")] == [38, 3, 2, 1]
        assert len([row for row in exceptions if row["source_row"] is not None]) == expected["row_exception_count"]
        assert len([row for row in exceptions if row["source_row"] is None]) == expected["file_exception_count"]
        assert runs[2]["source_row_count"] == expected["rejected_file_source_row_count"]
        assert sum(row["sale_amount"] for row in clean) == expected["total_sales_amount"]
        assert {summary.cell(row, 2).value: summary.cell(row, 3).value for row in (30, 31)} == {"B1": 26, "B2": 12}
        assert {summary.cell(row, 10).value: summary.cell(row, 11).value for row in (30, 31)} == {"Food": 33, "Home": 5}
        assert workbook["Exceptions"]["C4"].value == "FILE"
    finally:
        workbook.close()

    styled = load_workbook(output)
    try:
        assert len(styled["Summary"]._charts) == 2
        assert all(chart.dLbls.showCatName for chart in styled["Summary"]._charts)
        assert styled["Summary"]["B8"].number_format == "#,##0.00"
        assert "$" not in styled["Summary"]["B8"].number_format
        assert styled["Clean_Data"].freeze_panes == "D2"
        assert "CleanSales" in styled["Clean_Data"].tables
        assert styled["Clean_Data"].auto_filter.ref is None
        assert styled["Exceptions"].freeze_panes == "D2"
        assert styled["Run_Info"].auto_filter.ref is not None
    finally:
        styled.close()


def test_invalid_identity_and_invalid_row_do_not_claim_duplicate_key(tmp_path):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    write_csv(input_dir / "sales.csv", HEADERS, [
        ["2026-09-21", "B1", "O5", "", "P1", "Food", "1", "2", "COMPLETED"],
        ["2026-09-21", "B1", "O5", "1", "P1", "Food", "0", "2", "COMPLETED"],
        ["2026-09-21", "B1", "O5", "1", "P1", "Food", "1", "2", "COMPLETED"],
        ["2026-09-21", "B1", "O5", "1", "P1", "Food", "1", "2", "COMPLETED"],
        ["2026-09-21", "", "O6", "1", "P1", "Food", "1", "2", "COMPLETED"],
        ["2026-09-21", "B1", "O7", "1", "P1", "Food", "1", "2", "CANCELLED"],
    ])
    metrics = build_report(input_dir, tmp_path / "report.xlsx")
    assert metrics["clean_row_count"] == 1
    assert metrics["row_exception_count"] == 5
    assert metrics["duplicate_row_count"] == 1
    assert metrics["total_sales_amount"] == 2
    workbook = load_workbook(tmp_path / "report.xlsx", read_only=True, data_only=True)
    try:
        assert [row["error_codes"] for row in rows(workbook["Exceptions"])] == [
            "INVALID_LINE_NO", "INVALID_QUANTITY", "DUPLICATE_TRANSACTION",
            "MISSING_BRANCH_CODE", "UNSUPPORTED_STATUS",
        ]
    finally:
        workbook.close()


def test_output_cannot_overwrite_an_input_file(tmp_path):
    write_xlsx(tmp_path / "sales.xlsx", HEADERS, [
        ["2026-09-21", "B1", "O1", 1, "P1", "Food", 1, 2, "COMPLETED"],
    ])
    before = hash_file(tmp_path / "sales.xlsx")
    with pytest.raises(ValueError, match="outside the input directory"):
        build_report(tmp_path, tmp_path / "sales.xlsx")
    assert hash_file(tmp_path / "sales.xlsx") == before


def test_numeric_identifiers_accept_safe_integral_cells_and_reject_unsafe_ones(tmp_path):
    assert all(_identifier(value) is None for value in (float("nan"), float("inf"),
                                                        float("-inf"), 1.5, True))
    write_xlsx(tmp_path / "sales.xlsx", HEADERS, [
        ["2026-09-21", 100200, 8801234567890, 1, 100200.0, "Food", 1, 2, "COMPLETED"],
        ["2026-09-21", 1.5, "O2", 1, "P1", "Food", 1, 2, "COMPLETED"],
        ["2026-09-21", 10**16, "O3", 1, "P1", "Food", 1, 2, "COMPLETED"],
    ])
    output = tmp_path / "output" / "report.xlsx"
    metrics = build_report(tmp_path, output)
    assert (metrics["clean_row_count"], metrics["row_exception_count"], metrics["total_sales_amount"]) == (1, 2, 2)
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        assert [(row["branch_code"], row["order_id"], row["product_code"])
                for row in rows(workbook["Clean_Data"])] == [("100200", "8801234567890", "100200")]
        assert [row["error_codes"] for row in rows(workbook["Exceptions"])] == [
            "INVALID_BRANCH_CODE", "INVALID_BRANCH_CODE",
        ]
    finally:
        workbook.close()



def test_blank_residue_and_trailing_header_cells_are_ignored(tmp_path):
    source = tmp_path / "sales.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS + [None, None])
    sheet.append(["2026-09-21", "B1", "O1", 1, "P1", "Food", 1, 2, "COMPLETED"])
    sheet["M50"].fill = PatternFill("solid", fgColor="FFFF00")
    workbook.save(source)
    workbook.close()
    metrics = build_report(tmp_path, tmp_path / "out" / "report.xlsx")
    assert (metrics["total_source_row_count"], metrics["clean_row_count"], metrics["row_exception_count"]) == (1, 1, 0)


def test_excel_saved_styled_residue_fixture(tmp_path):
    source = Path(__file__).parent / "fixtures" / "excel_styled_residue.xlsx"
    shutil.copy2(source, tmp_path / source.name)
    metrics = build_report(tmp_path, tmp_path / "out" / "report.xlsx")
    assert (metrics["total_source_row_count"], metrics["clean_row_count"], metrics["row_exception_count"]) == (1, 1, 0)


def test_stale_xlsx_dimension_does_not_hide_real_rows(tmp_path):
    original = tmp_path / "original.xlsx"
    write_xlsx(original, HEADERS, [["2026-09-21", "B1", "O1", 1, "P1", "Food", 1, 2, "COMPLETED"]])
    with zipfile.ZipFile(original) as archive:
        with zipfile.ZipFile(tmp_path / "sales.xlsx", "w") as stale:
            for member in archive.namelist():
                payload = archive.read(member)
                if member == "xl/worksheets/sheet1.xml":
                    assert b'<dimension ref="A1:I2"' in payload
                    payload = payload.replace(b'<dimension ref="A1:I2"', b'<dimension ref="A1:A1"')
                stale.writestr(member, payload)
    original.unlink()
    metrics = build_report(tmp_path, tmp_path / "out" / "report.xlsx")
    assert (metrics["total_source_row_count"], metrics["clean_row_count"]) == (1, 1)


def test_csv_blank_lines_and_duplicate_nonempty_header(tmp_path):
    source = tmp_path / "sales.csv"
    write_csv(source, HEADERS + ["", ""], [
        ["2026-09-21", "B1", "O1", 1, "P1", "Food", 1, 2, "COMPLETED", "", ""],
        [], ["", ""],
    ])
    output = tmp_path / "out" / "report.xlsx"
    assert build_report(tmp_path, output)["clean_row_count"] == 1
    assert build_report(tmp_path, output)["total_source_row_count"] == 1
    write_csv(source, HEADERS + [" ORDER_ID ", ""], [])
    metrics = build_report(tmp_path, output)
    assert metrics["rejected_file_count"] == 1
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        assert "Duplicate column names" in rows(workbook["Exceptions"])[0]["error_detail"]
    finally:
        workbook.close()


def test_title_row_reports_missing_columns(tmp_path):
    write_csv(tmp_path / "sales.csv", ["Weekly sales export", ""], [])
    output = tmp_path / "out" / "report.xlsx"
    build_report(tmp_path, output)
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        assert rows(workbook["Exceptions"])[0]["error_codes"] == "MISSING_REQUIRED_COLUMNS"
    finally:
        workbook.close()


def test_malformed_xlsx_and_unsupported_spreadsheets_are_isolated(tmp_path):
    write_csv(tmp_path / "good.csv", HEADERS, [
        ["2026-09-21", "B1", "O1", 1, "P1", "Food", 1, 2, "COMPLETED"],
    ])
    good_xlsx = tmp_path / "good.xlsx"
    write_xlsx(good_xlsx, HEADERS, [])
    with zipfile.ZipFile(good_xlsx) as archive:
        with zipfile.ZipFile(tmp_path / "bad_xml.xlsx", "w") as damaged:
            for member in archive.namelist():
                payload = b"<worksheet><broken>" if member == "xl/worksheets/sheet1.xml" else archive.read(member)
                damaged.writestr(member, payload)
    (tmp_path / "bad_zip.xlsx").write_bytes(b"not a zip file")
    with zipfile.ZipFile(good_xlsx) as archive:
        with zipfile.ZipFile(tmp_path / "bad_stream.xlsx", "w", compression=zipfile.ZIP_STORED) as damaged:
            for member in archive.namelist():
                damaged.writestr(member, archive.read(member))
    damaged_bytes = (tmp_path / "bad_stream.xlsx").read_bytes()
    assert b"<worksheet " in damaged_bytes
    (tmp_path / "bad_stream.xlsx").write_bytes(damaged_bytes.replace(b"<worksheet ", b"<Worksheet ", 1))
    (tmp_path / "old.xls").write_bytes(b"unsupported")
    (tmp_path / "macro.xlsm").write_bytes(b"unsupported")
    (tmp_path / "~$lock.xlsx").write_bytes(b"lock")
    output = tmp_path / "out" / "report.xlsx"
    metrics = build_report(tmp_path, output)
    assert (metrics["clean_row_count"], metrics["rejected_file_count"], metrics["discovered_input_file_count"]) == (1, 5, 7)
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        failures = {row["source_file"]: row["error_codes"] for row in rows(workbook["Exceptions"])}
        assert failures == {"bad_xml.xlsx": "FILE_READ_ERROR", "bad_zip.xlsx": "FILE_READ_ERROR",
                            "bad_stream.xlsx": "FILE_READ_ERROR",
                            "macro.xlsm": "UNSUPPORTED_FILE_FORMAT", "old.xls": "UNSUPPORTED_FILE_FORMAT"}
        assert all(row["source_file"] != "~$lock.xlsx" for row in rows(workbook["Exceptions"]))
    finally:
        workbook.close()


def test_illegal_customer_control_is_visible_and_other_rows_survive(tmp_path):
    write_csv(tmp_path / "sales.csv", HEADERS, [
        ["2026-09-21", "B1", "O1", 1, "P1", "Food\x01bad", 1, 2, "COMPLETED"],
        ["2026-09-21", "B1", "O2", 1, "P1", "Food", 1, 3, "COMPLETED"],
    ])
    output = tmp_path / "out" / "report.xlsx"
    metrics = build_report(tmp_path, output)
    assert (metrics["clean_row_count"], metrics["row_exception_count"]) == (1, 1)
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        rejected = rows(workbook["Exceptions"])[0]
        assert rejected["error_codes"] == "EXCEL_ILLEGAL_CHARACTER"
        assert rejected["original_category"] == "Food\\x01bad"
        assert rows(workbook["Clean_Data"])[0]["order_id"] == "O2"
    finally:
        workbook.close()


def test_cp949_korean_and_formula_like_text_round_trip_as_literal(tmp_path):
    source = tmp_path / "sales.csv"
    with source.open("w", encoding="cp949", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(HEADERS)
        writer.writerow(["2026-09-21", "=1+1", "O1", 1, "P1", "한글 식품", 1, 2, "COMPLETED"])
    output = tmp_path / "out" / "report.xlsx"
    assert build_report(tmp_path, output)["clean_row_count"] == 1
    workbook = load_workbook(output)
    try:
        cell = workbook["Clean_Data"]["B2"]
        assert cell.value == "=1+1" and cell.data_type == "s"
        assert workbook["Clean_Data"]["F2"].value == "한글 식품"
        assert workbook["Summary"]["B30"].data_type == "s"
    finally:
        workbook.close()


def test_csv_bytes_undecodable_by_both_attempted_encodings_are_file_failure(tmp_path):
    (tmp_path / "bad.csv").write_bytes(b"\xff\xff\xff")
    write_csv(tmp_path / "good.csv", HEADERS, [["2026-09-21", "B1", "O1", 1, "P1", "Food", 1, 2, "COMPLETED"]])
    output = tmp_path / "out" / "report.xlsx"
    assert build_report(tmp_path, output)["clean_row_count"] == 1
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        failure = rows(workbook["Exceptions"])[0]
        assert failure["error_codes"] == "FILE_READ_ERROR"
        assert "UTF-8" in failure["error_detail"] and "CP949" in failure["error_detail"]
    finally:
        workbook.close()


def test_conflicting_duplicate_names_changed_fields_without_changing_first(tmp_path):
    write_csv(tmp_path / "sales.csv", HEADERS, [
        ["2026-09-21", "B1", "O1", 1, "P1", "Food", 1, 2, "COMPLETED"],
        ["2026-09-21", "B2", "O1", 1, "P1", "Food", 2, 2, "COMPLETED"],
    ])
    output = tmp_path / "out" / "report.xlsx"
    metrics = build_report(tmp_path, output)
    assert metrics["total_sales_amount"] == 2
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        detail = rows(workbook["Exceptions"])[0]["error_detail"]
        assert "branch_code" in detail and "quantity" in detail and "sales.csv:2" in detail
    finally:
        workbook.close()


@pytest.mark.parametrize("excel_fixture", [False, True], ids=["openpyxl", "excel-saved-xml"])
def test_sparse_xlsx_trailing_status_is_row_exception_and_other_data_survives(tmp_path, excel_fixture):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    source = input_dir / "01_sparse.xlsx"
    if excel_fixture:
        fixture = Path(__file__).parent / "fixtures" / "excel_styled_residue.xlsx"
        with zipfile.ZipFile(fixture) as original, zipfile.ZipFile(source, "w") as sparse:
            for name in original.namelist():
                payload = original.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    cell = b'<c r="I2" t="s"><v>13</v></c>'
                    assert cell in payload
                    payload = payload.replace(cell, b"")
                sparse.writestr(name, payload)
    else:
        write_xlsx(source, HEADERS, [
            ["2026-09-21", "B1", "O1", 1, "P1", "Food", 1, 2],
            ["2026-09-21", "B1", "O2", 1, "P1", "Food", 1, 3, "COMPLETED"],
        ])
    write_csv(input_dir / "02_good.csv", HEADERS, [
        ["2026-09-21", "B2", "O3", 1, "P1", "Food", 2, 4, "COMPLETED"],
    ])
    before = {path.name: hash_file(path) for path in input_dir.iterdir()}
    output = tmp_path / "report.xlsx"
    metrics = build_report(input_dir, output)
    assert metrics["total_source_row_count"] == (2 if excel_fixture else 3)
    assert metrics["clean_row_count"] == (1 if excel_fixture else 2)
    assert metrics["row_exception_count"] == 1
    assert metrics["rejected_file_count"] == 0
    assert metrics["total_sales_amount"] == (8 if excel_fixture else 11)
    assert metrics["total_source_row_count"] == (
        metrics["clean_row_count"] + metrics["row_exception_count"]
        + metrics["rejected_file_source_row_count"]
    )
    assert {path.name: hash_file(path) for path in input_dir.iterdir()} == before
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        exception, = rows(workbook["Exceptions"])
        assert exception["error_codes"] == "MISSING_STATUS"
        assert (exception["source_file"], exception["source_row"], exception["failure_scope"]) == (
            "01_sparse.xlsx", 2, "ROW",
        )
        assert exception["original_status"] is None
        assert exception["original_quantity"] == 1
        assert rows(workbook["Run_Info"])[0]["exception_contribution"] == 1
    finally:
        workbook.close()


@pytest.mark.parametrize("field", HEADERS + ["extra_note", "all"])
@pytest.mark.parametrize("cached", [False, True], ids=["no-cache", "stale-cache"])
def test_xlsx_formula_rows_are_visible_rejections_regardless_of_cache(tmp_path, field, cached):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    headers = HEADERS + ["extra_note"]
    values = ["2026-09-21", "B1", "O1", 1, "P1", "Food", 1, 2, "COMPLETED", "note"]
    affected = list(range(len(headers))) if field == "all" else [headers.index(field)]
    formula_row = values.copy()
    for index in affected:
        formula_row[index] = "=1+1"
    source = input_dir / "01_formula.xlsx"
    write_xlsx(source, headers, [formula_row, values])
    if cached:
        original = tmp_path / "original.xlsx"
        source.rename(original)
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with zipfile.ZipFile(original) as archive, zipfile.ZipFile(source, "w") as stale:
            for name in archive.namelist():
                payload = archive.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    xml = ET.fromstring(payload)
                    for cell in xml.findall(".//s:c", ns):
                        if cell.find("s:f", ns) is not None:
                            assert cell.find("s:f", ns).text == "1+1"
                            cell.find("s:v", ns).text = "999"
                    payload = ET.tostring(xml)
                stale.writestr(name, payload)
        probe = load_workbook(source, read_only=True, data_only=True)
        try:
            assert probe.active.cell(2, affected[0] + 1).value == 999
        finally:
            probe.close()
    write_csv(input_dir / "02_good.csv", HEADERS, [
        ["2026-09-21", "B2", "O2", 1, "P1", "Food", 2, 3, "COMPLETED"],
        ["2026-09-21", "B1", "O1", 1, "P1", "Food", 1, 2, "COMPLETED"],
    ])
    before = {path.name: hash_file(path) for path in input_dir.iterdir()}
    output = tmp_path / "report.xlsx"
    metrics = build_report(input_dir, output)
    assert metrics["total_source_row_count"] == 4
    assert metrics["clean_row_count"] == 2
    assert metrics["row_exception_count"] == 2
    assert metrics["duplicate_row_count"] == 1
    assert metrics["rejected_file_count"] == 0
    assert metrics["rejected_file_source_row_count"] == 0
    assert metrics["total_sales_amount"] == 8
    assert {path.name: hash_file(path) for path in input_dir.iterdir()} == before
    workbook = load_workbook(output, data_only=False)
    try:
        rejected, duplicate = rows(workbook["Exceptions"])
        assert rejected["error_codes"] == "UNSUPPORTED_FORMULA"
        assert (rejected["source_file"], rejected["source_row"], rejected["failure_scope"]) == (
            "01_formula.xlsx", 2, "ROW",
        )
        for index in affected:
            name = headers[index]
            assert name in rejected["error_detail"]
            assert rejected["original_" + name] == "=1+1"
            column = list(workbook["Exceptions"].values)[0].index("original_" + name) + 1
            assert workbook["Exceptions"].cell(2, column).data_type == "s"
        assert duplicate["error_codes"] == "DUPLICATE_TRANSACTION"
        assert [row["source_file"] for row in rows(workbook["Clean_Data"])] == [
            "01_formula.xlsx", "02_good.csv",
        ]
        assert rows(workbook["Run_Info"])[0]["exception_contribution"] == 1
    finally:
        workbook.close()


@pytest.mark.parametrize("formula", [
    ArrayFormula(ref="H2", text="=1+1"), DataTableFormula(ref="H2", r1="H1"),
], ids=["array", "data-table"])
def test_structured_xlsx_formulas_are_inspectable_exceptions(tmp_path, formula):
    source = tmp_path / "sales.xlsx"
    write_xlsx(source, HEADERS, [
        ["2026-09-21", "B1", "O1", 1, "P1", "Food", 1, formula, "COMPLETED"],
    ])
    before = hash_file(source)
    output = tmp_path / "out" / "report.xlsx"
    metrics = build_report(tmp_path, output)
    assert (metrics["total_source_row_count"], metrics["clean_row_count"], metrics["row_exception_count"]) == (1, 0, 1)
    assert hash_file(source) == before
    workbook = load_workbook(output)
    try:
        exception, = rows(workbook["Exceptions"])
        assert exception["error_codes"] == "UNSUPPORTED_FORMULA"
        assert "unit_price" in exception["error_detail"]
        if isinstance(formula, ArrayFormula):
            assert exception["original_unit_price"] == "=1+1"
        else:
            assert "DataTableFormula" in exception["original_unit_price"]
            assert "H2" in exception["original_unit_price"]
    finally:
        workbook.close()


def test_xlsx_formula_like_literal_text_is_accepted_as_text(tmp_path):
    source = tmp_path / "sales.xlsx"
    workbook = Workbook()
    workbook.active.append(HEADERS)
    workbook.active.append(["2026-09-21", "=1+1", "O1", 1, "P1", "Food", 1, 2, "COMPLETED"])
    workbook.active["B2"].data_type = "s"
    workbook.save(source)
    workbook.close()
    output = tmp_path / "out" / "report.xlsx"
    assert build_report(tmp_path, output)["clean_row_count"] == 1
    result = load_workbook(output)
    try:
        assert result["Clean_Data"]["B2"].value == "=1+1"
        assert result["Clean_Data"]["B2"].data_type == "s"
        assert result["Summary"]["B30"].data_type == "s"
    finally:
        result.close()
